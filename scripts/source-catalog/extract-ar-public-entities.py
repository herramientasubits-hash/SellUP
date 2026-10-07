#!/usr/bin/env python3
"""
AR — directorio de entidades públicas (municipios, organismos nacionales,
universidades nacionales) con CUIT, web y tamaño → un CSV normalizado
(SOURCES-AR-PUBLIC-ENTITIES-1).

Lee archivos LOCALES (los baja la persona que corre la carga; todos CC-BY 4.0 o
públicos) y escribe `cuit,name,entity_type,province,website,size,size_kind,match`.
NUNCA guarda autoridades, DNI, CUIL, teléfonos, direcciones ni correos.

  · ReFeGLo (Ministerio del Interior, gobiernos locales, web + población 2022):
      http://datos.mininterior.gob.ar/dataset/8431ab79-3f44-41ee-b72b-cdc645f68d1e/resource/ebe52783-6a7f-4112-8a7f-a17c91cb6b4e/download/brgl_-portal-andino-23_12.csv
  · Padrón de ARCA (CUIT + denominación de 30 caracteres), descomprimido:
      https://www.afip.gob.ar/genericos/cInscripcion/archivos/apellidoNombreDenominacion.zip
  · SIPRO (COMPR.AR, «Organismo Publico» con CUIT y provincia):
      https://infra.datos.gob.ar/catalog/jgm/dataset/4/distribution/4.23/download/SiPRO.csv
  · INDEC, dotación de la APN, empresas y sociedades del Estado (XLSX, hoja más reciente):
      https://www.indec.gob.ar/ftp/cuadros/economia/serie_dotacion_apn_cuadro_2_3.xlsx
  · Mapa del Estado 2019 (sólo la columna `web`; se ignoran las de autoridades):
      http://infra.datos.gob.ar/catalog/jgm/dataset/2/distribution/2.1/download/estructura-20191209.csv
  · Wikidata, universidades de Argentina con sitio oficial (P856), JSON de SPARQL.

Reglas (medidas el 07-10-2026):
  · Municipio: ≥ --min-population habitantes (censo 2022). CUIT por nombre contra
    el padrón: «MUNICIPALIDAD DE <gobierno_local>» (abreviaturas GRAL/CNEL/STA…
    expandidas). Vale si hay UN solo CUIT (exacto, o prefijo cuando el padrón
    corta el nombre a 30), o si entre varios el SIPRO confirma uno en la MISMA
    provincia. Un parecido parcial (el padrón omite palabras) sólo vale con SIPRO.
  · Organismo nacional: dotación INDEC ≥ --min-workers. CUIT por nombre exacto
    (o primeros 30) en el padrón, UNO solo. Web del Mapa del Estado sólo si el
    dominio lleva la sigla o una palabra distintiva del nombre; nunca
    argentina.gob.ar (compartida por muchas entidades).
  · Universidad nacional: las «Organismo Publico» del SIPRO cuyo nombre lleva
    UNIVERSIDAD; web de Wikidata por nombre.

Uso:
  python3 scripts/source-catalog/extract-ar-public-entities.py \
    --refeglo=brgl.csv --padron=SELE-SAL-CONSTA.p20out1.<fecha>.tmp --sipro=SiPRO.csv \
    --indec=serie_dotacion_apn_cuadro_2_3.xlsx --mapa=estructura-20191209.csv \
    --wikidata=wd_univ.json --out=ar-public-entities.csv

Requiere: pip install openpyxl
"""

import argparse
import collections
import csv
import json
import re
import unicodedata

import openpyxl

ABBR = {
    'GRAL': 'GENERAL', 'CNEL': 'CORONEL', 'PTE': 'PRESIDENTE', 'PDTE': 'PRESIDENTE', 'STA': 'SANTA',
    'STO': 'SANTO', 'TTE': 'TENIENTE', 'CMTE': 'COMANDANTE', 'COMTE': 'COMANDANTE', 'DR': 'DOCTOR',
    'GOB': 'GOBERNADOR', 'ING': 'INGENIERO', 'PBRO': 'PRESBITERO',
}
SHARED_DOMAINS = {'argentina.gob.ar', 'gob.ar', 'gov.ar', 'facebook.com', 'instagram.com', 'twitter.com', 'x.com'}
NAME_STOP = {'DE', 'DEL', 'LA', 'LAS', 'LOS', 'EL', 'Y', 'NACIONAL', 'INSTITUTO', 'ENTE', 'AGENCIA', 'COMISION',
             'SERVICIO', 'ADMINISTRACION', 'MINISTERIO', 'SECRETARIA', 'DIRECCION', 'REGISTRO', 'CONSEJO'}
MUNI_LINE = re.compile(r'^3[034]\d{9}(MUNICIPALIDAD|MUNICIPIO)')


def norm(s):
    s = unicodedata.normalize('NFD', s or '').encode('ascii', 'ignore').decode().upper()
    s = re.sub(r'\([^)]*\)', ' ', s)
    s = re.sub(r'[^A-Z0-9 ]', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()


def expand(s):
    return ' '.join(ABBR.get(w, w) for w in s.split())


def muni_core(name):
    n = norm(name)
    n = re.sub(r'^(MUNICIPALIDAD|MUNICIPIO)( DE)? ', '', n)
    n = re.sub(r'^(LA CIUDAD DE|CIUDAD DE|DEL PARTIDO DE|PARTIDO DE) ', '', n)
    return expand(n)


def host(url):
    h = (url or '').strip().lower()
    h = re.sub(r'^[a-z]+://', '', h).split('/')[0].split('?')[0].split(':')[0]
    h = re.sub(r'^www\d?\.', '', h)
    return h if re.match(r'^[a-z0-9-]+(\.[a-z0-9-]+)+$', h) else ''


def registrable(h):
    parts = h.split('.')
    if len(parts) >= 3 and len(parts[-1]) == 2 and parts[-2] in {'com', 'gob', 'gov', 'org', 'edu', 'net', 'mil', 'int', 'tur'}:
        return '.'.join(parts[-3:])
    return '.'.join(parts[-2:])


def read_padron(path):
    """CUIT → denominación (30), sólo personas jurídicas."""
    muni, all_names = [], collections.defaultdict(set)
    with open(path, 'rb') as f:
        for raw in f:
            s = raw.decode('latin-1')
            cuit = s[:11]
            if cuit[:2] not in ('30', '33', '34') or not cuit.isdigit():
                continue
            name = s[11:41].strip()
            all_names[norm(name)].add(cuit)
            if MUNI_LINE.match(s):
                muni.append({'cuit': cuit, 'name': name, 'core': muni_core(name), 'trunc': len(name) >= 29})
    return muni, all_names


def read_sipro(path):
    out = {}
    for r in csv.DictReader(open(path, encoding='utf-8', errors='replace')):
        if (r.get('Tipo_de_Personeria') or '').strip().lower().startswith('organismo'):
            cuit = re.sub(r'\D', '', r.get('CUIT_NIT') or '')
            if len(cuit) == 11 and cuit[:2] in ('30', '33', '34'):
                out[cuit] = {'name': (r.get('Razon _Social') or r.get('Razon_Social') or '').strip(), 'prov': norm(r.get('Provincia'))}
    return out


def match_municipality(local, prov, muni, sipro):
    k, p = expand(norm(local)), norm(prov)
    tiers = [
        ('exact', [m for m in muni if m['core'] == k]),
        ('truncated', [m for m in muni if m['trunc'] and len(m['core']) >= 8 and k.startswith(m['core'])]),
        ('partial', [m for m in muni if len(m['core']) >= 5 and (k.startswith(m['core'] + ' ') or k.endswith(' ' + m['core']))]),
    ]
    for kind, found in tiers:
        cuits = {m['cuit']: m for m in found}
        if not cuits:
            continue
        confirmed = [c for c in cuits if sipro.get(c, {}).get('prov') == p]
        if kind != 'partial' and len(cuits) == 1:
            return next(iter(cuits)), kind
        if len(confirmed) == 1:
            return confirmed[0], kind + '+sipro'
        return None, kind + ':ambiguous'
    return None, 'none'


def read_indec(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for row in wb[wb.sheetnames[-1]].iter_rows(values_only=True):
        name, size = row[0], row[1]
        if not isinstance(name, str) or not isinstance(size, (int, float)):
            continue
        if name.startswith(('Total', 'Administración', 'Empresas', 'Otros entes')):
            continue
        sigla = re.findall(r'\(([A-ZÁÉÍÓÚ]{2,})\)', name)
        out.append({'name': re.sub(r'\s*\(\d+\)\s*', ' ', name).strip(), 'size': int(size), 'sigla': sigla[0].lower() if sigla else ''})
    return out


def read_mapa_webs(path):
    """nombre de unidad normalizado → host, sin columnas de autoridades."""
    out = {}
    for r in csv.DictReader(open(path, encoding='utf-8', errors='replace')):
        unit = (r.get('unidad') or '').encode('latin-1', 'ignore').decode('utf-8', 'ignore') or r.get('unidad') or ''
        h = host(r.get('web'))
        if h and norm(unit) not in out:
            out[norm(unit)] = h
    return out


def web_fits_name(h, name, sigla):
    if not h or registrable(h) in SHARED_DOMAINS:
        return False
    label = registrable(h).split('.')[0]
    if sigla and (label == sigla or label.startswith(sigla)):
        return True
    words = [w.lower() for w in norm(name).split() if w not in NAME_STOP and len(w) >= 5]
    return any(w in label for w in words)


def main():
    ap = argparse.ArgumentParser()
    for a in ('refeglo', 'padron', 'sipro', 'indec', 'mapa', 'wikidata', 'out'):
        ap.add_argument('--' + a, required=True)
    ap.add_argument('--min-population', type=int, default=20000)
    ap.add_argument('--min-workers', type=int, default=200)
    args = ap.parse_args()

    muni, padron_names = read_padron(args.padron)
    sipro = read_sipro(args.sipro)
    rows, stats = [], collections.Counter()

    # Municipios.
    for r in csv.DictReader(open(args.refeglo, encoding='cp1252', errors='replace'), delimiter=';'):
        pop = int(re.sub(r'\D', '', r.get('población_censo_2022') or '') or 0)
        if pop < args.min_population:
            continue
        cuit, how = match_municipality(r['gobierno_local'], r['provincia'], muni, sipro)
        stats['muni:' + how.split(':')[0] + (':amb' if 'ambiguous' in how else '')] += 1
        if not cuit:
            continue
        h = host(r.get('sitio_web'))
        rows.append([cuit, 'Municipalidad de ' + r['gobierno_local'].strip(), 'municipality', r['provincia'].strip(),
                     h if h and registrable(h) not in SHARED_DOMAINS else '', pop, 'population_2022', how])

    # Organismos nacionales.
    mapa = read_mapa_webs(args.mapa)
    for e in read_indec(args.indec):
        if e['size'] < args.min_workers:
            continue
        k = norm(e['name'])
        cuits = padron_names.get(k) or padron_names.get(k[:30].strip()) or set()
        stats['national:' + ('unique' if len(cuits) == 1 else 'none' if not cuits else 'ambiguous')] += 1
        if len(cuits) != 1:
            continue
        h = mapa.get(k, '')
        rows.append([next(iter(cuits)), e['name'], 'national_entity', 'Nacional',
                     h if web_fits_name(h, e['name'], e['sigla']) else '', e['size'], 'workers_indec', 'exact'])

    # Universidades nacionales (SIPRO) + web de Wikidata.
    wd = {}
    for b in json.load(open(args.wikidata))['results']['bindings']:
        wd.setdefault(norm(b['uLabel']['value']), host(b['web']['value']))
    for cuit, s in sipro.items():
        if 'UNIVERSIDAD' not in norm(s['name']):
            continue
        h = wd.get(norm(s['name']), '')
        stats['university:' + ('web' if h else 'no_web')] += 1
        rows.append([cuit, s['name'], 'university', s['prov'].title(), h, '', '', 'sipro'])

    seen, out = set(), []
    for row in rows:
        if row[0] in seen:
            continue
        seen.add(row[0])
        out.append(row)
    with open(args.out, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['cuit', 'name', 'entity_type', 'province', 'website', 'size', 'size_kind', 'match'])
        w.writerows(out)
    print(json.dumps({'rows': len(out), **stats}, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()
