#!/usr/bin/env python3
"""
MX — listas públicas con RFC de persona moral → un CSV normalizado
(SOURCES-MX-RFC-PUBLIC-LISTS-1).

Lee archivos LOCALES (los baja la persona que corre la carga) y escribe
`rfc,name,list,extra`. No guarda teléfonos, correos, domicilios ni representantes.

  · SAT Padrón de Importadores (PDF):
      https://www.sat.gob.mx/minisitio/PadronImportadoresExportadores/documentos/Pad_Imp.pdf
    y sus sectoriales Pad_Imp_Sec.pdf / Pad_Exp_Sec.pdf (misma carpeta).
  · SAT Directorio de Donatarias Autorizadas (XLSX), sólo «Activo» / «Reactivado»:
      https://www.sat.gob.mx/minisitio/DonatariasAutorizadas/documentos/padron_donatarias/Dir_2026.xlsx
    extra = tipo de donataria (letra del SAT: A asistencial, B educativa…).
  · Nuevo León, padrón de proveedores (XLSX); extra = estratificación:
      https://www.nl.gob.mx/ (Secretaría de Administración → padrón de proveedores)
  · Ciudad de México, SAF, padrón de proveedores de transparencia (XLSX, art. 121
    fr. XXXIV, uno por año); extra = estratificación:
      https://transparencia.finanzas.cdmx.gob.mx/transparencia/public/articulos_121_XXXIV
  · Ciudad de México, datos abiertos «padrón de proveedores vigente» (CSV, morales):
      https://datos.cdmx.gob.mx/dataset/padron-de-proveedores-vigente

El ETL (`run-mx-rfc-public-lists-etl.ts`) es quien filtra persona moral, normaliza
y decide la fila; aquí sólo se extrae.

Uso:
  python3 scripts/source-catalog/extract-mx-rfc-public-lists.py \
    --importadores=Pad_Imp.pdf --importadores-sec=Pad_Imp_Sec.pdf --exportadores-sec=Pad_Exp_Sec.pdf \
    --donatarias=Dir_2026.xlsx --nuevo-leon=padron_proveedores.xlsx \
    --cdmx-saf A121Fr34_2021.xlsx … A121Fr34_2026.xlsx --cdmx-datos proveedores_datos_abiertos.csv \
    --out=mx-rfc-lists.csv

Requiere: pip install pymupdf openpyxl
"""

import argparse
import csv
import re

import fitz  # PyMuPDF
import openpyxl

RFC = re.compile(r'^[A-ZÑ&]{3}\d{6}[A-Z0-9]{3}$')
ACTIVE_DONATARIA = {'ACTIVO', 'REACTIVADO'}


def from_pdf(path, tag):
    # Cada registro del PDF sale como líneas «ID», «RFC», «NOMBRE».
    out = []
    for page in fitz.open(path):
        lines = [l.strip() for l in page.get_text().split('\n') if l.strip()]
        for i, line in enumerate(lines[:-1]):
            if RFC.match(line):
                out.append((line, lines[i + 1], tag, ''))
    return out


def header_index(rows):
    for i, row in enumerate(rows):
        cells = [str(c).strip().upper() if c else '' for c in (row or [])]
        if 'RFC' in cells:
            return i, cells
    raise SystemExit('sin cabecera con columna RFC')


def column(header, *needles):
    for i, name in enumerate(header):
        if any(n in name for n in needles):
            return i
    return None


def from_donatarias(path):
    ws = openpyxl.load_workbook(path, read_only=True).active
    ws.reset_dimensions()
    rows = list(ws.iter_rows(values_only=True))
    at, header = header_index(rows)
    rfc, name = header.index('RFC'), column(header, 'DENOMINACI', 'RAZ')
    kind, status = column(header, 'ACTIVIDAD'), column(header, 'ESTADO DEL CONTRIBUYENTE')
    out = []
    for row in rows[at + 1:]:
        if not row or not row[rfc]:
            continue
        if status is not None and str(row[status] or '').strip().upper() not in ACTIVE_DONATARIA:
            continue
        out.append((str(row[rfc]).strip().upper(), str(row[name] or '').strip(), 'sat_donatarias',
                    str(row[kind] or '').strip()[:4] if kind is not None else ''))
    return out


def from_nuevo_leon(path):
    ws = openpyxl.load_workbook(path, read_only=True).active
    ws.reset_dimensions()
    rows = list(ws.iter_rows(values_only=True))
    at, header = header_index(rows)
    rfc, name, size = header.index('RFC'), column(header, 'RAZ'), column(header, 'ESTRATIF')
    out = []
    for row in rows[at + 1:]:
        if not row or not row[rfc]:
            continue
        out.append((str(row[rfc]).strip().upper().replace(' ', '').replace('-', ''), str(row[name] or '').strip(),
                    'nl_proveedores', str(row[size] or '').strip().upper() if size is not None else ''))
    return out


def from_cdmx_saf(paths):
    # Ciudad de México, SAF: padrón de proveedores de transparencia (art. 121 fr.
    # XXXIV, plantilla SIPOT). Sólo personas morales (columna de razón social). Se
    # devuelven del año más NUEVO al más viejo: el ETL se queda con el primer tamaño.
    out = []
    dated = []
    for path in paths:
        ws = openpyxl.load_workbook(path, read_only=True).active
        ws.reset_dimensions()
        rows = list(ws.iter_rows(values_only=True))
        at = next(i for i, r in enumerate(rows[:15]) if r and any(isinstance(c, str) and 'RFC' in c.upper() for c in r))
        header = [str(c or '').strip().upper() for c in rows[at]]
        rfc = next(i for i, h in enumerate(header) if 'REGISTRO FEDERAL DE CONTRIBUYENTES' in h or h.startswith('RFC'))
        name = column(header, 'DENOMINACI')
        size = column(header, 'ESTRATIFIC')
        year_col = column(header, 'EJERCICIO')
        for row in rows[at + 1:]:
            if not row or not row[rfc] or name is None or not row[name]:
                continue
            year = row[year_col] if year_col is not None else 0
            dated.append((int(year or 0), (str(row[rfc]).strip().upper().replace(' ', '').replace('-', ''),
                                           str(row[name]).strip(), 'cdmx_proveedores',
                                           str(row[size] or '').strip() if size is not None else '')))
    dated.sort(key=lambda item: -item[0])
    out += [item[1] for item in dated]
    return out


def from_cdmx_datos(path):
    # Ciudad de México, datos abiertos «padrón de proveedores vigente» (CSV): sólo morales.
    out = []
    with open(path, newline='', encoding='utf-8') as f:
        for row in csv.DictReader(f):
            if (row.get('tipo_persona') or '').strip().upper() != 'MORAL':
                continue
            out.append(((row.get('RFC') or '').strip().upper(), (row.get('razon_social') or '').strip(), 'cdmx_proveedores_datos', ''))
    return out


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--importadores')
    p.add_argument('--importadores-sec')
    p.add_argument('--exportadores-sec')
    p.add_argument('--donatarias')
    p.add_argument('--nuevo-leon')
    p.add_argument('--cdmx-saf', nargs='*', default=[])
    p.add_argument('--cdmx-datos')
    p.add_argument('--out', required=True)
    a = p.parse_args()
    out = []
    if a.importadores:
        out += from_pdf(a.importadores, 'sat_importadores')
    if a.importadores_sec:
        out += from_pdf(a.importadores_sec, 'sat_importadores_sectorial')
    if a.exportadores_sec:
        out += from_pdf(a.exportadores_sec, 'sat_exportadores_sectorial')
    if a.donatarias:
        out += from_donatarias(a.donatarias)
    if a.nuevo_leon:
        out += from_nuevo_leon(a.nuevo_leon)
    if a.cdmx_saf:
        out += from_cdmx_saf(a.cdmx_saf)
    if a.cdmx_datos:
        out += from_cdmx_datos(a.cdmx_datos)
    with open(a.out, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['rfc', 'name', 'list', 'extra'])
        w.writerows(out)
    print(f'{len(out)} filas → {a.out}')


if __name__ == '__main__':
    main()
